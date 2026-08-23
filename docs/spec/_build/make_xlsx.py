# -*- coding: utf-8 -*-
"""CookPilot_기능명세서_HIPO.xlsx 를 만든다."""
import os
import re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from hipo_data import SYSTEM, MODULES, FEATURES
from hipo_funcs import FUNCS
from hipo_tables import CHECKS, DATA, SCREENS

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   "CookPilot_기능명세서_HIPO.xlsx")

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="맑은 고딕", size=9)
ID_FONT = Font(name="맑은 고딕", size=9, bold=True)
MONO_FONT = Font(name="Consolas", size=8.5)
LV_FILL = {0: PatternFill("solid", fgColor="D6DCE4"),
           1: PatternFill("solid", fgColor="E9EDF4"),
           2: PatternFill("solid", fgColor="F5F7FB")}
THIN = Side(style="thin", color="B4B4B4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP_WRAP = Alignment(vertical="top", wrap_text=True)
TOP_CENTER = Alignment(vertical="top", horizontal="center", wrap_text=True)


def sheet(wb, title, headers, rows, widths, mono_cols=(), id_cols=(0,),
          center_cols=(), row_fill=None):
    """헤더 · 자동 필터 · 틀 고정 · 열 너비 · 줄바꿈을 적용한 시트를 만든다."""
    ws = wb.create_sheet(title)
    ws.append(list(headers))
    for c, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = TOP_CENTER
        cell.border = BORDER
    for r, row in enumerate(rows, start=2):
        ws.append(list(row))
        fill = row_fill(r - 2, row) if row_fill else None
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = (MONO_FONT if (c - 1) in mono_cols
                         else ID_FONT if (c - 1) in id_cols else BODY_FONT)
            cell.alignment = TOP_CENTER if (c - 1) in center_cols else TOP_WRAP
            cell.border = BORDER
            if fill:
                cell.fill = fill
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)),
                                      max(1, len(rows) + 1))
    ws.sheet_view.zoomScale = 100
    return ws


def evidence(loc):
    """구현위치 '경로 :: 함수() :F#' → 근거 '경로:F#'."""
    m = re.match(r"^(.+?) :: .+? :F(\d+)$", loc)
    return "%s:F%s" % (m.group(1), m.group(2)) if m else loc


MOD_NAME = {m[0]: m[1] for m in MODULES}
MOD_PURPOSE = {m[0]: m[2] for m in MODULES}
FEAT = {f[0]: f for f in FEATURES}

wb = Workbook()
wb.remove(wb.active)

# ── 01_모듈기능명세서 ────────────────────────────────────
rows = []
for f in FEATURES:
    (fn, name, mid, desc, i, pr, o, pre, post, ext, data, scr,
     funcs, loc, ev) = f
    rows.append([mid, MOD_NAME[mid], MOD_PURPOSE[mid], fn, name, desc,
                 i, pr, o, pre, post, ext, data, scr, funcs, loc, ev])
rows.sort(key=lambda r: (r[0], r[3]))
sheet(wb, "01_모듈기능명세서",
      ["모듈ID", "모듈명", "모듈목적", "기능ID", "기능명", "기능설명", "입력", "처리",
       "출력", "선행조건", "후행조건", "외부연계", "관련데이터", "관련화면", "관련함수",
       "구현위치", "근거"],
      rows,
      [8, 14, 34, 8, 22, 40, 26, 46, 30, 22, 30, 18, 20, 12, 32, 46, 40],
      mono_cols=(15, 16), id_cols=(0, 3), center_cols=(0, 3, 12, 13))

# ── 02_기능체크리스트 ────────────────────────────────────
sheet(wb, "02_기능체크리스트",
      ["체크ID", "모듈ID", "기능ID", "체크항목", "예상동작", "구현상태", "구현위치",
       "관련함수", "테스트필요", "비고", "근거"],
      [list(c) for c in CHECKS],
      [8, 8, 8, 44, 40, 10, 46, 30, 9, 34, 38],
      mono_cols=(6, 10), id_cols=(0, 1, 2), center_cols=(0, 1, 2, 5, 8))

# ── 03_HIPO ─────────────────────────────────────────────
hipo = [[0, SYSTEM[0], "-", "시스템", SYSTEM[1], SYSTEM[2],
         "사용자 입력 · 브라우저 저장 값 · 외부 API 응답",
         "레시피 생성 · 조리 진행 · 게시글 · 상담을 모듈 단위로 처리",
         "화면 렌더링 · 브라우저 저장소 기록 · Supabase 반영", "-",
         "D1 · D2 · D3 · D4", "J1 – J7"]]
for mid, mname, mpurpose in MODULES:
    fns = [f for f in FEATURES if f[2] == mid]
    hipo.append([1, mid, SYSTEM[0], "모듈", mname, mpurpose,
                 " / ".join("%s: %s" % (f[0], f[4]) for f in fns),
                 "소속 기능 %s 를 수행한다" % " · ".join(f[0] for f in fns),
                 " / ".join("%s: %s" % (f[0], f[6]) for f in fns), "-",
                 " · ".join(sorted({d for f in fns for d in f[10].split(" · ")})),
                 " · ".join(sorted({s for f in fns for s in f[11].split(" · ")}))])
    for f in fns:
        hipo.append([2, f[0], mid, "기능", f[1], f[3], f[4], f[5], f[6],
                     f[13], f[10], f[11]])
        for fu in [x for x in FUNCS if x[3] == f[0]]:
            hipo.append([3, fu[0], f[0], "함수", fu[1],
                         "%s(%s) 기능의 구현 함수" % (f[1], f[0]),
                         fu[4], fu[5], fu[6], fu[10], f[10], f[11]])
sheet(wb, "03_HIPO",
      ["계층번호", "HIPO_ID", "상위ID", "유형", "이름", "목적", "입력", "처리", "출력",
       "구현위치", "관련데이터", "관련화면"],
      hipo,
      [8, 11, 9, 8, 26, 44, 30, 44, 30, 46, 20, 12],
      mono_cols=(9,), id_cols=(1, 2), center_cols=(0, 1, 2, 3, 10, 11),
      row_fill=lambda idx, row: LV_FILL.get(row[0]))

# ── 04_함수IPO ──────────────────────────────────────────
sheet(wb, "04_함수IPO",
      ["함수ID", "함수명", "모듈ID", "기능ID", "입력", "처리", "출력", "호출자",
       "호출대상", "외부I/O", "구현위치", "근거"],
      [list(f) for f in FUNCS],
      [10, 26, 8, 8, 30, 46, 30, 28, 30, 24, 48, 36],
      mono_cols=(1, 10, 11), id_cols=(0, 2, 3), center_cols=(0, 2, 3))

# ── 05_추적성매트릭스 ───────────────────────────────────
trace = []
for fu in FUNCS:
    f = FEAT[fu[3]]
    trace.append([fu[2], fu[3], fu[0], evidence(fu[10]), f[10], f[11],
                  SYSTEM[0], fu[9] if fu[9] != "없음" else ""])
sheet(wb, "05_추적성매트릭스",
      ["모듈ID", "기능ID", "함수ID", "코드근거", "데이터ID", "화면ID", "시스템ID", "비고"],
      trace,
      [8, 8, 10, 52, 22, 14, 9, 30],
      mono_cols=(3,), id_cols=(0, 1, 2), center_cols=(0, 1, 2, 4, 5, 6))

# ── 06_데이터사전 ───────────────────────────────────────
sheet(wb, "06_데이터사전",
      ["데이터ID", "저장소", "키/테이블/버킷", "데이터명", "설명", "값/내용",
       "생성/쓰기 위치", "읽기 위치", "삭제 위치", "관련모듈", "관련기능", "근거"],
      [list(d) for d in DATA],
      [10, 16, 30, 20, 40, 30, 42, 42, 42, 10, 40, 40],
      mono_cols=(2, 6, 7, 8, 11), id_cols=(0,), center_cols=(0, 1, 9))

# ── 07_화면여정 ─────────────────────────────────────────
sheet(wb, "07_화면여정",
      ["화면ID", "Route", "화면명", "진입조건", "입력데이터", "주요기능", "다음화면",
       "전달데이터", "분기조건", "이전/복귀경로", "관련기능", "근거"],
      [list(s) for s in SCREENS],
      [10, 22, 20, 20, 26, 32, 40, 18, 40, 34, 18, 36],
      mono_cols=(1, 6, 11), id_cols=(0,), center_cols=(0, 10))

wb.save(OUT)
print("wrote", OUT)
print("시트:", ", ".join(wb.sheetnames))
print("행 수 - 01:%d 02:%d 03:%d 04:%d 05:%d 06:%d 07:%d"
      % (len(FEATURES), len(CHECKS), len(hipo), len(FUNCS), len(trace),
         len(DATA), len(SCREENS)))
